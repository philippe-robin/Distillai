"""Tests du pipeline : parseurs, extraction sur fixture, export Excel."""

import json
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from logic_immo import extract  # noqa: E402
from logic_immo.export_xlsx import write_workbook  # noqa: E402

FIXTURE = Path(__file__).parent / "fixture_results.html"


def check(label, condition, detail=""):
    status = "OK  " if condition else "FAIL"
    print(f"[{status}] {label}{(' — ' + str(detail)) if detail and not condition else ''}")
    return bool(condition)


def test_parsers():
    ok = True
    ok &= check("prix '745 000 €'", extract.to_number("745 000 €") == 745000)
    ok &= check("prix insécable", extract.to_number("1 250 000 €") == 1250000)
    ok &= check("prix '550.000'", extract.to_number("550.000") == 550000)
    ok &= check("prix dict", extract.to_number({"value": 612000}) == 612000)
    ok &= check("surface '186,5 m²'", extract.to_number("186,5") == 186.5)
    card = extract.parse_card_text("Maison 7 pièces 4 chambres\n178 m²\n67380 Lingolsheim\n675 000 €")
    ok &= check("carte -> prix", card.get("prix") == 675000, card)
    ok &= check("carte -> surface", card.get("surface") == 178, card)
    ok &= check("carte -> pièces", card.get("pieces") == 7, card)
    ok &= check("carte -> chambres", card.get("chambres") == 4, card)
    ok &= check("carte -> ville", card.get("ville") == "Lingolsheim", card)
    ok &= check("carte -> CP", card.get("cp") == "67380", card)
    ok &= check("bruit ignoré", extract.parse_card_text("Estimez votre bien").get("prix") is None)
    return ok


def test_json_harvest():
    payload = json.loads(
        FIXTURE.read_text(encoding="utf-8").split('type="application/json">')[1].split("</script>")[0]
    )
    found = extract.harvest_json(payload, source="api")
    ok = check("2 annonces JSON", len(found) == 2, len(found))
    by_id = {r["id"]: r for r in found}
    a = by_id.get("LI-1001", {})
    ok &= check("prix imbriqué", a.get("prix") == 789000, a.get("prix"))
    ok &= check("surface", a.get("surface") == 232, a.get("surface"))
    ok &= check("terrain", a.get("terrain") == 812, a.get("terrain"))
    ok &= check("ville sous-objet", a.get("ville") == "Mundolsheim", a.get("ville"))
    ok &= check("agence", a.get("agence") == "Agence Alpha", a.get("agence"))
    ok &= check("URL absolue", str(a.get("url")).startswith("https://www.logic-immo.com/"), a.get("url"))
    b = by_id.get("LI-1002", {})
    ok &= check("isNew -> Neuf", b.get("type_projet") == "Neuf", b.get("type_projet"))
    ok &= check("advertiser -> agence", b.get("agence") == "Beta Promotion", b.get("agence"))
    return ok


def test_dedupe_and_filter():
    records = [
        {"id": "A", "prix": 700000, "surface": 200, "ville": "X", "agence": None},
        {"id": "A", "prix": 700000, "surface": 200, "ville": "X", "agence": "Comble le vide"},
        {"id": None, "url": "https://x/1?utm=a", "prix": 600000, "surface": 160},
        {"id": None, "url": "https://x/1?utm=b", "prix": 600000, "surface": 160},
    ]
    deduped = extract.dedupe(records)
    ok = check("dédup id + URL", len(deduped) == 2, len(deduped))
    cross = extract.dedupe([
        {"id": "LI-9", "url": "https://x/maison-li-9", "prix": 700000, "surface": 200, "agence": "Alpha"},
        {"id": None, "url": "https://x/maison-li-9/", "prix": 700000, "surface": 200, "dpe": "C"},
    ])
    ok &= check("dédup JSON + DOM", len(cross) == 1, cross)
    ok &= check("fusion inter-sources", cross[0].get("agence") == "Alpha" and cross[0].get("dpe") == "C", cross)
    ok &= check("fusion des champs", deduped[0]["agence"] == "Comble le vide")
    criteria = {"priceMin": 550000, "priceMax": 800000, "spaceMin": 150, "spaceMax": 250}
    ok &= check("dans critères", extract.matches_criteria({"prix": 700000, "surface": 200}, criteria))
    ok &= check("prix trop haut", not extract.matches_criteria({"prix": 900000, "surface": 200}, criteria))
    ok &= check("surface trop faible", not extract.matches_criteria({"prix": 700000, "surface": 90}, criteria))
    ok &= check("champ absent toléré", extract.matches_criteria({"prix": 700000}, criteria))
    return ok


def test_browser_pipeline():
    """Charge la fixture dans Chromium : valide next_data + DOM + dédup inter-sources."""
    from scrape_logic_immo import _harvest_page

    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("[SKIP] Playwright non installé")
        return True
    with sync_playwright() as pw:
        try:
            browser = pw.chromium.launch(headless=True)
        except Exception:
            browser = pw.chromium.launch(headless=True, executable_path="/opt/pw-browsers/chromium")
        page = browser.new_page()
        page.goto(FIXTURE.as_uri())
        found = extract.dedupe(_harvest_page(page))
        browser.close()
    ok = check("3 annonces uniques (JSON + DOM)", len(found) == 3, [f.get("id") or f.get("url") for f in found])
    ok &= check("pub écartée", all("Estimez" not in (f.get("titre") or "") for f in found))
    ok &= check(
        "annonce DOM complète",
        any(f.get("prix") == 675000 and f.get("surface") == 178 for f in found),
        found,
    )
    return ok


def test_workbook():
    records = [
        {"titre": "Maison A", "ville": "Mundolsheim", "cp": "67450", "prix": 789000, "surface": 232,
         "terrain": 812, "pieces": 8, "chambres": 5, "dpe": "B", "type_projet": "Ancien",
         "agence": "Agence Alpha", "url": "https://www.logic-immo.com/a"},
        {"titre": "Maison B", "ville": "Oberhausbergen", "cp": "67205", "prix": 612000, "surface": 168,
         "pieces": 6, "chambres": 4, "type_projet": "Neuf", "agence": "Beta", "url": None},
    ]
    out = Path(tempfile.gettempdir()) / "test_biens.xlsx"
    write_workbook(records, str(out), criteria={"priceMin": 550000, "zone": "Strasbourg"}, source_url="https://x")
    from openpyxl import load_workbook

    wb = load_workbook(out)
    ws = wb["Biens"]
    ok = check("feuilles", wb.sheetnames == ["Biens", "Critères"], wb.sheetnames)
    ok &= check("2 lignes", ws.max_row == 3, ws.max_row)
    # A : 789000/232 = 3401 €/m² ; B : 612000/168 = 3643 €/m² -> A en premier.
    ok &= check("tri prix/m² croissant", ws["B2"].value == "Maison A", ws["B2"].value)
    ok &= check("formule prix/m²", str(ws["J2"].value).startswith("=IFERROR("), ws["J2"].value)
    ok &= check("filtres actifs", ws.auto_filter.ref is not None)
    ok &= check("volets figés", ws.freeze_panes == "B2")
    ok &= check("URL vide tolérée", ws["N3"].value is None or ws["N3"].value == "")
    print(f"       classeur de test : {out}")
    return ok


if __name__ == "__main__":
    results = [
        ("parseurs", test_parsers()),
        ("extraction JSON", test_json_harvest()),
        ("dédup / filtre", test_dedupe_and_filter()),
        ("pipeline navigateur", test_browser_pipeline()),
        ("classeur Excel", test_workbook()),
    ]
    failed = [name for name, ok in results if not ok]
    print("\n" + ("Tous les tests passent." if not failed else f"Échecs : {', '.join(failed)}"))
    raise SystemExit(1 if failed else 0)
